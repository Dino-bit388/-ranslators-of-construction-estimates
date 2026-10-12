import re
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from models.enums import EstimateRowType
from models.header_info import HeaderInfo
from models.section import Section
from models.work import Work
from services.parser.normalization import normalize_number
from services.parser.normalization import normalize_text
from services.parser.normalization import normalize_unit
from services.parser.header_patterns import NORMATIVE_CODE_PATTERN, SECTION_PATTERN, TOTAL_PATTERN
from models.estimate import Estimate
from models.resource import Resource
from services.parser.header_patterns import _match_header_field



class ExcelParser:
    from services.parser.header_patterns import find_header

    def __init__(self, filepath: str):
        self.filepath = filepath
        self.workbook = None                          
        self.merged_ranges: list = []                  
        self.header_info: HeaderInfo | None = None      
        self.column_map: dict[str, int] = {}             
        self.current_section: Section | None = None       
        self.current_work: Work | None = None
        self.diagnostics: list[dict] = []

    def _looks_like_data_row(self, row: list) -> bool:
        text = normalize_text(row[0] if row else None)
        if NORMATIVE_CODE_PATTERN.search(text):
            return True
        has_text = any(isinstance(v, str) and v.strip() for v in row)
        has_number = any(isinstance(v, (int, float)) for v in row)
        return has_text and has_number

    def find_header(self, ws: Worksheet, max_scan_rows: int = 40) -> HeaderInfo | None:
        columns: dict[str, int] = {}
        start_row: int | None = None
        end_row: int | None = None

        for row_number in range(1, max_scan_rows + 1):
            row = [ws.cell(row=row_number, column=c).value for c in range(1, ws.max_column + 1)]

            found_in_this_row = {}
            for col_index, cell_value in enumerate(row, start=1):
                if cell_value is None:
                    continue
                text = normalize_text(cell_value)
                field_name = _match_header_field(text)
                if field_name and field_name not in columns:
                    found_in_this_row[field_name] = col_index

            if found_in_this_row:
                columns.update(found_in_this_row)
                if start_row is None:
                    start_row = row_number
                end_row = row_number
                continue

            
            if columns and self._looks_like_data_row(row):
                data_start_row = row_number
                return HeaderInfo(
                    start_row=start_row,
                    end_row=end_row,
                    data_start_row=data_start_row,
                    columns=columns,
                )

        return None

    def _validate_file(self) -> None:
        if not self.filepath.lower().endswith((".xlsx", ".xlsm")):
            raise ValueError(f"Неподдерживаемый формат файла: {self.filepath}")

    def load(self) -> None:
        self._validate_file()
        self.workbook = load_workbook(self.filepath, data_only=True)

    def parse(self) -> "Estimate":
        from models.estimate import Estimate  

        self.load()

        for sheet_name in self.workbook.sheetnames:
            ws = self.workbook[sheet_name]
            self.merged_ranges = list(ws.merged_cells.ranges)

            header_info = self.find_header(ws)
            if header_info is None:
                self.diagnostics.append({"sheet": sheet_name, "status": "header_not_found"})
                continue

            self.header_info = header_info
            self.column_map = header_info.columns

            metadata = self.parse_metadata(ws, header_info)
            sections = self._parse_rows(ws, header_info)

            return Estimate(metadata=metadata, sections=sections)

        raise ValueError("Не удалось найти таблицу материалов ни на одном листе файла")

    def parse_metadata(self, ws: Worksheet, header_info: HeaderInfo):
        from models.metadata import EstimateMetadata

        object_name = None
        estimate_name = None

        for row_number in range(1, header_info.start_row):
            for col in range(1, ws.max_column + 1):
                value = ws.cell(row=row_number, column=col).value
                if not value:
                        continue
                text = str(value)
                if "объект" in text.lower() and object_name is None:
                    object_name = text
                if "смета" in text.lower() and estimate_name is None:
                    estimate_name = text

        return EstimateMetadata(object_name=object_name, estimate_name=estimate_name)

    def _get_row_style(self, ws: Worksheet, row_number: int) -> dict:
        is_merged_wide = any(
            rng.min_row <= row_number <= rng.max_row and (rng.max_col - rng.min_col) >= 3
            for rng in self.merged_ranges
        )
        is_bold = any(
            ws.cell(row=row_number, column=c).font and ws.cell(row=row_number, column=c).font.bold
            for c in range(1, ws.max_column + 1)
        )
        filled_count = sum(
            1 for c in range(1, ws.max_column + 1)
            if ws.cell(row=row_number, column=c).value is not None
        )
        return {"is_merged_wide": is_merged_wide, "is_bold": is_bold, "filled_count": filled_count}

    def detect_row_type(self, row: list, row_number: int, ws: Worksheet) -> tuple[EstimateRowType, float]:
        name_col = self.column_map.get("name")
        unit_col = self.column_map.get("unit")
        quantity_col = self.column_map.get("quantity")

        text = normalize_text(row[name_col - 1]) if name_col else ""
        unit_value = row[unit_col - 1] if unit_col else None
        quantity_value = row[quantity_col - 1] if quantity_col else None

        style = self._get_row_style(ws, row_number)

        scores = {t: 0 for t in EstimateRowType}

        if style["is_merged_wide"]:
            scores[EstimateRowType.SECTION] += 2
            scores[EstimateRowType.TOTAL] += 1
        if style["is_bold"]:
            scores[EstimateRowType.SECTION] += 1
            scores[EstimateRowType.TOTAL] += 1
        if style["filled_count"] <= 2:
            scores[EstimateRowType.SECTION] += 1
            scores[EstimateRowType.TOTAL] += 1

        if SECTION_PATTERN.search(text):
            scores[EstimateRowType.SECTION] += 3
        if TOTAL_PATTERN.search(text):
            scores[EstimateRowType.TOTAL] += 3
        if NORMATIVE_CODE_PATTERN.search(text):
            scores[EstimateRowType.WORK] += 3
        if text and unit_value and quantity_value is not None:
            scores[EstimateRowType.RESOURCE] += 2

        best_type = max(scores, key=scores.get)
        best_score = scores[best_type]
        total_score = sum(scores.values())

        if best_score == 0:
            return EstimateRowType.UNKNOWN, 0.0

        confidence = best_score / total_score if total_score else 0.0
        return best_type, confidence

    def extract_section(self, row: list, row_number: int) -> Section:
        name = normalize_text(row[self.column_map["name"] - 1])
        return Section(name=name, number=row_number)

    def extract_work(self, row: list, row_number: int) -> Work:
        code_col = self.column_map.get("code")
        code = normalize_text(row[code_col - 1]) if code_col else None
        name = normalize_text(row[self.column_map["name"] - 1])
        unit = self._get_unit(row)
        quantity = self._get_quantity(row)
        return Work(code=code or None, name=name, unit=unit, quantity=quantity)

    def extract_resource(self, row: list, row_number: int) -> "Resource":
        from models.resource import Resource

        name = normalize_text(row[self.column_map["name"] - 1])
        unit = self._get_unit(row)
        quantity = self._get_quantity(row)
        return Resource(name=name, unit=unit, quantity=quantity, row_number=row_number)

    def _get_unit(self, row: list) -> str | None:
        col = self.column_map.get("unit")
        return normalize_unit(row[col - 1]) if col else None

    def _get_quantity(self, row: list) -> float | None:
        col = self.column_map.get("quantity")
        return normalize_number(row[col - 1]) if col else None  

    def _parse_rows(self, ws: Worksheet, header_info: HeaderInfo) -> list[Section]:
        sections: list[Section] = []
        self.current_section = None
        self.current_work = None

        for row_number, row in enumerate(
            ws.iter_rows(min_row=header_info.data_start_row, values_only=True),
            start=header_info.data_start_row,
        ):
            row = list(row)
            if not any(value is not None for value in row):
                continue  

            row_type, confidence = self.detect_row_type(row, row_number, ws)
            self.diagnostics.append({
                "row": row_number, "type": row_type.value, "confidence": round(confidence, 2),
            })

            if row_type == EstimateRowType.SECTION:
                self.current_section = self.extract_section(row, row_number)
                sections.append(self.current_section)
                self.current_work = None

            elif row_type == EstimateRowType.WORK:
                if self.current_section is None:
                    self.current_section = Section(name="Без раздела")
                    sections.append(self.current_section)
                self.current_work = self.extract_work(row, row_number)
                self.current_section.works.append(self.current_work)

            elif row_type == EstimateRowType.RESOURCE:
                resource = self.extract_resource(row, row_number)
                self._attach_resource(sections, resource)

            elif row_type == EstimateRowType.TOTAL:
                continue  

            else:   
                resource = self.extract_resource(row, row_number)
                resource.name = f"[?] {resource.name}"
                self._attach_resource(sections, resource)

        return sections

    def _attach_resource(self, sections: list[Section], resource) -> None:
        """Кладёт ресурс в текущую работу; если работы нет — создаёт служебную"""
        if self.current_section is None:
            self.current_section = Section(name="Без раздела")
            sections.append(self.current_section)
        if self.current_work is None:
            self.current_work = Work(code=None, name="—", unit=None, quantity=None)
            self.current_section.works.append(self.current_work)
        self.current_work.resources.append(resource)               